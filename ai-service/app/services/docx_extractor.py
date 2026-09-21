"""
DOCX text extraction using python-docx.
Extracts text from paragraphs and tables.
"""

import io
import logging
from dataclasses import dataclass

from docx import Document

logger = logging.getLogger(__name__)


@dataclass
class DocxExtractionResult:
    """Result of DOCX text extraction."""

    text: str
    paragraph_count: int
    table_count: int
    method: str = "python-docx"


def extract_text_from_docx(file_bytes: bytes) -> DocxExtractionResult:
    """Extract text from a DOCX file.

    Strategy:
    - Extracts all paragraph text in document order.
    - Extracts table content formatted as readable text.
    - Preserves heading structure with markdown-like formatting.

    Args:
        file_bytes: Raw DOCX content.

    Returns:
        DocxExtractionResult with extracted text and counts.
    """
    doc = Document(io.BytesIO(file_bytes))
    text_parts: list[str] = []

    # --- Extract paragraphs ---
    paragraph_count = 0
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        paragraph_count += 1

        # Preserve heading structure
        if para.style and para.style.name.startswith("Heading"):
            level = para.style.name.replace("Heading", "").strip()
            try:
                level_int = int(level)
                prefix = "#" * min(level_int, 3)
                text_parts.append(f"{prefix} {text}")
            except ValueError:
                text_parts.append(text)
        else:
            text_parts.append(text)

    # --- Extract tables ---
    table_count = len(doc.tables)
    for table_idx, table in enumerate(doc.tables):
        table_rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            # Skip empty rows
            if any(cells):
                table_rows.append(" | ".join(cells))

        if table_rows:
            text_parts.append("")  # Blank line before table
            text_parts.extend(table_rows)
            text_parts.append("")  # Blank line after table

    full_text = "\n".join(text_parts)

    logger.info(
        "DOCX extracted: %d paragraphs, %d tables, %d chars",
        paragraph_count,
        table_count,
        len(full_text),
    )

    return DocxExtractionResult(
        text=full_text,
        paragraph_count=paragraph_count,
        table_count=table_count,
    )
